#!/usr/bin/env python3
"""Convert full PDFs into page-located Chinese DOCX files.

The script supports text PDFs and scanned/image PDFs. Text pages are extracted
with pdfplumber. Pages with little/no embedded text are rendered with
ImageMagick and OCRed with Tesseract. English text can be translated to Chinese
with an offline Argos model or online fallback services.

Usage:
  python scripts/pdf_to_chinese_docx.py input.pdf output.docx --translate auto
  python scripts/pdf_to_chinese_docx.py input_dir output_dir --batch --translate argos
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import shlex
import subprocess
import sys
import tempfile
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Iterable

import pdfplumber
import requests
from docx import Document


def run(cmd: list[str], env: dict[str, str] | None = None) -> subprocess.CompletedProcess[str]:
    return subprocess.run(cmd, text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env)


def command_exists(name: str) -> bool:
    return shutil.which(name) is not None


def clean_text(text: str) -> str:
    text = text.replace("\u00a0", " ")
    text = re.sub(r"([a-z])([A-Z])", r"\1 \2", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def is_mostly_english(text: str) -> bool:
    latin = len(re.findall(r"[A-Za-z]", text))
    cjk = len(re.findall(r"[\u4e00-\u9fff]", text))
    return latin >= 80 and latin > cjk * 2


def split_for_translation(text: str, max_chars: int = 1200) -> list[str]:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    for para in paragraphs:
        if len(para) <= max_chars:
            chunks.append(para)
            continue
        sentences = re.split(r"(?<=[.;:?!])\s+", para)
        buf = ""
        for sent in sentences:
            if len(buf) + len(sent) + 1 <= max_chars:
                buf = (buf + " " + sent).strip()
            else:
                if buf:
                    chunks.append(buf)
                buf = sent
        if buf:
            chunks.append(buf)
    return chunks


def load_cache(path: Path) -> dict[str, str]:
    if path.exists():
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def save_cache(path: Path, cache: dict[str, str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")


class Translator:
    def __init__(self, backend: str, cache_path: Path, sleep: float = 0.4):
        self.backend = backend
        self.cache_path = cache_path
        self.cache = load_cache(cache_path)
        self.sleep = sleep
        self._argos_translate = None

    def translate(self, text: str) -> tuple[str, str]:
        text = clean_text(text)
        if not text:
            return "", "empty"
        if not is_mostly_english(text):
            return text, "source-is-chinese-or-mixed"

        translated_chunks: list[str] = []
        methods: list[str] = []
        for chunk in split_for_translation(text):
            key = hashlib.sha256((self.backend + "\n" + chunk).encode("utf-8")).hexdigest()
            if key in self.cache:
                translated_chunks.append(self.cache[key])
                methods.append("cache")
                continue
            translated, method = self._translate_chunk(chunk)
            self.cache[key] = translated
            translated_chunks.append(translated)
            methods.append(method)
            if method.startswith("mymemory"):
                time.sleep(self.sleep)
        save_cache(self.cache_path, self.cache)
        method_summary = ",".join(sorted(set(methods)))
        return "\n\n".join(translated_chunks), method_summary

    def _translate_chunk(self, chunk: str) -> tuple[str, str]:
        backends: list[str]
        if self.backend == "auto":
            backends = ["argos", "mymemory", "none"]
        else:
            backends = [self.backend]
        for backend in backends:
            if backend == "argos":
                translated = self._translate_argos(chunk)
                if translated:
                    return translated, "argos"
            elif backend == "mymemory":
                translated = self._translate_mymemory(chunk)
                if translated:
                    return translated, "mymemory"
            elif backend == "none":
                return "[未翻译] " + chunk, "none"
        return "[未翻译] " + chunk, "none"

    def _translate_argos(self, chunk: str) -> str | None:
        try:
            if self._argos_translate is None:
                import argostranslate.translate  # type: ignore

                self._argos_translate = argostranslate.translate
            translated = self._argos_translate.translate(chunk, "en", "zh")
            translated = clean_text(translated)
            if translated and translated != chunk:
                return translated
        except Exception:
            return None
        return None

    def _translate_mymemory(self, chunk: str) -> str | None:
        try:
            resp = requests.get(
                "https://api.mymemory.translated.net/get",
                params={"q": chunk[:4900], "langpair": "en|zh-CN"},
                timeout=30,
            )
            if resp.status_code != 200:
                return None
            data = resp.json()
            translated = clean_text(data.get("responseData", {}).get("translatedText", ""))
            if translated and translated.lower() != chunk.lower():
                return translated
        except Exception:
            return None
        return None


def tesseract_langs(cmd: list[str], env: dict[str, str] | None = None) -> set[str]:
    if not cmd:
        return set()
    cp = run(cmd + ["--list-langs"], env=env)
    lines = (cp.stdout + "\n" + cp.stderr).splitlines()
    return {line.strip() for line in lines if line.strip() and not line.lower().startswith("list of available")}


def choose_tesseract_cmd(lang: str, preferred: str) -> tuple[list[str], dict[str, str] | None, set[str], list[str]]:
    required = {p for p in lang.replace("+", " ").split() if p}
    warnings: list[str] = []
    candidates: list[tuple[list[str], dict[str, str] | None]] = []
    if preferred != "auto":
        candidates.append((shlex.split(preferred), None))
    else:
        if command_exists("tesseract"):
            candidates.append((["tesseract"], None))
            conda_tessdata = Path("/Users/xiluo/miniconda3/envs/claude/share/tessdata")
            if conda_tessdata.exists():
                env = os.environ.copy()
                env["TESSDATA_PREFIX"] = str(conda_tessdata)
                candidates.append((["tesseract"], env))
        if command_exists("conda"):
            candidates.append((["conda", "run", "-n", "claude", "tesseract"], None))

    best_cmd: list[str] = []
    best_env: dict[str, str] | None = None
    best_langs: set[str] = set()
    for cmd, env in candidates:
        langs = tesseract_langs(cmd, env=env)
        if langs and not best_cmd:
            best_cmd, best_env, best_langs = cmd, env, langs
        if required <= langs:
            return cmd, env, langs, warnings
    if best_cmd:
        missing = required - best_langs
        if missing:
            warnings.append(
                f"Selected Tesseract command lacks languages {', '.join(sorted(missing))}: {' '.join(best_cmd)}"
            )
        return best_cmd, best_env, best_langs, warnings
    warnings.append("No usable Tesseract command found")
    return [], None, set(), warnings


def render_page_with_magick(pdf_path: Path, page_index: int, out_png: Path, dpi: int) -> None:
    if not command_exists("magick"):
        raise RuntimeError("ImageMagick 'magick' command not found")
    cmd = [
        "magick",
        "-density",
        str(dpi),
        f"{pdf_path}[{page_index}]",
        "-background",
        "white",
        "-alpha",
        "remove",
        "-alpha",
        "off",
        "-colorspace",
        "Gray",
        "-strip",
        str(out_png),
    ]
    cp = run(cmd)
    if cp.returncode != 0 or not out_png.exists():
        raise RuntimeError((cp.stderr or cp.stdout).strip())


def ocr_image(image_path: Path, lang: str, psm: int, tesseract_cmd: list[str], tesseract_env: dict[str, str] | None) -> str:
    cp = run(tesseract_cmd + [str(image_path), "stdout", "-l", lang, "--psm", str(psm)], env=tesseract_env)
    if cp.returncode != 0:
        raise RuntimeError((cp.stderr or cp.stdout).strip())
    return clean_text(cp.stdout)


def ocr_pdf_page(
    pdf_path: Path,
    page_index: int,
    work_dir: Path,
    dpi: int,
    lang: str,
    psm: int,
    tesseract_cmd: list[str],
    tesseract_env: dict[str, str] | None,
) -> tuple[int, str]:
    image_path = work_dir / f"page_{page_index + 1:04d}.png"
    render_page_with_magick(pdf_path, page_index, image_path, dpi)
    return page_index, ocr_image(image_path, lang, psm, tesseract_cmd, tesseract_env)


def extract_pdf_pages(pdf_path: Path, args: argparse.Namespace) -> tuple[list[dict], list[str]]:
    warnings: list[str] = []
    tesseract_cmd, tesseract_env, langs, tess_warnings = choose_tesseract_cmd(args.ocr_lang, args.tesseract_cmd)
    warnings.extend(tess_warnings)
    required = {p for p in args.ocr_lang.replace("+", " ").split() if p}
    missing = required - langs

    pages: list[dict] = []
    ocr_tasks: list[int] = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        total = len(pdf.pages)
        limit = total if args.max_pages <= 0 else min(total, args.max_pages)
        with tempfile.TemporaryDirectory(prefix="pdf_docx_ocr_") as temp:
            temp_dir = Path(temp)
            for idx in range(limit):
                page = pdf.pages[idx]
                embedded = clean_text(page.extract_text(x_tolerance=2, y_tolerance=3) or "")
                method = "embedded_text"
                text = embedded
                if args.force_ocr or len(embedded) < args.min_text_chars:
                    method = "ocr"
                    if missing:
                        text = embedded
                        method = "ocr-skipped-missing-language"
                    else:
                        text = ""
                        ocr_tasks.append(idx)
                pages.append({
                    "source_file": pdf_path.name,
                    "page_no": idx + 1,
                    "method": method,
                    "text": text,
                    "char_count": len(text),
                })

            if ocr_tasks:
                max_workers = max(1, args.ocr_workers)
                with ThreadPoolExecutor(max_workers=max_workers) as executor:
                    futures = {
                        executor.submit(
                            ocr_pdf_page,
                            pdf_path,
                            idx,
                            temp_dir,
                            args.dpi,
                            args.ocr_lang,
                            args.psm,
                            tesseract_cmd,
                            tesseract_env,
                        ): idx
                        for idx in ocr_tasks
                    }
                    for future in as_completed(futures):
                        idx = futures[future]
                        try:
                            _, text = future.result()
                            pages[idx]["text"] = text
                            pages[idx]["char_count"] = len(text)
                        except Exception as exc:
                            warnings.append(f"{pdf_path.name} page={idx + 1} OCR failed: {exc}")
                            pages[idx]["text"] = ""
                            pages[idx]["char_count"] = 0
                            pages[idx]["method"] = "ocr-error"
    return pages, warnings


def add_page_to_doc(doc: Document, page: dict, translated: str, translation_method: str, include_original: bool) -> None:
    heading = f"{page['source_file']} | page={page['page_no']} | {page['method']}"
    doc.add_heading(heading, level=2)
    if translated:
        doc.add_paragraph(translated)
    else:
        doc.add_paragraph("[空白页或未识别出文字]")
    if include_original and page["text"] and page["text"] != translated:
        doc.add_paragraph("原文：")
        doc.add_paragraph(page["text"])
    doc.add_paragraph(f"translation_method={translation_method}; char_count={page['char_count']}")


def convert_one(pdf_path: Path, output_docx: Path, args: argparse.Namespace, cache_dir: Path) -> dict:
    pages, warnings = extract_pdf_pages(pdf_path, args)
    translator = Translator(args.translate, cache_dir / "translation_cache.json", sleep=args.sleep)

    doc = Document()
    doc.add_heading(pdf_path.stem, level=1)
    doc.add_paragraph(f"来源文件：{pdf_path.name}")
    doc.add_paragraph(f"生成说明：按页解析 PDF；扫描页走 OCR；英文段落按 translate={args.translate} 转为中文。")

    translated_pages = 0
    ocr_pages = 0
    methods: dict[str, int] = {}
    for page_idx, page in enumerate(pages, start=1):
        if args.progress and (page_idx == 1 or page_idx % args.progress_every == 0 or page_idx == len(pages)):
            print(f"[{pdf_path.name}] translating/writing page {page_idx}/{len(pages)}", flush=True)
        translated, method = translator.translate(page["text"])
        add_page_to_doc(doc, page, translated, method, args.include_original)
        translated_pages += 1 if method not in {"empty", "source-is-chinese-or-mixed", "none"} else 0
        ocr_pages += 1 if page["method"].startswith("ocr") else 0
        methods[method] = methods.get(method, 0) + 1

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_docx))
    manifest = {
        "source_file": pdf_path.name,
        "output_docx": str(output_docx),
        "page_count": len(pages),
        "ocr_pages": ocr_pages,
        "translated_pages": translated_pages,
        "translation_methods": methods,
        "warnings": warnings,
    }
    (output_docx.with_suffix(".manifest.json")).write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def iter_pdfs(input_path: Path) -> Iterable[Path]:
    if input_path.is_file() and input_path.suffix.lower() == ".pdf":
        yield input_path
    elif input_path.is_dir():
        yield from sorted(input_path.rglob("*.pdf"))
    else:
        raise FileNotFoundError(f"No PDF input found: {input_path}")


def safe_stem(path: Path) -> str:
    return re.sub(r"[\\/:*?\"<>|]+", "_", path.stem).strip() or "document"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Convert text/scanned PDFs into page-located Chinese DOCX files.")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--batch", action="store_true", help="Treat output as a directory and convert every PDF under input")
    parser.add_argument("--translate", choices=["auto", "argos", "mymemory", "none"], default="auto")
    parser.add_argument("--ocr-lang", default="chi_sim+eng")
    parser.add_argument(
        "--tesseract-cmd",
        default="auto",
        help="Tesseract command. Use 'auto' to prefer a command with requested languages, including conda run -n claude tesseract.",
    )
    parser.add_argument("--dpi", type=int, default=220)
    parser.add_argument("--psm", type=int, default=6)
    parser.add_argument("--ocr-workers", type=int, default=2, help="Parallel OCR workers for scanned pages")
    parser.add_argument("--min-text-chars", type=int, default=30)
    parser.add_argument("--max-pages", type=int, default=0)
    parser.add_argument("--force-ocr", action="store_true")
    parser.add_argument("--include-original", action="store_true")
    parser.add_argument("--progress", action="store_true", help="Print per-file progress")
    parser.add_argument("--progress-every", type=int, default=25)
    parser.add_argument("--sleep", type=float, default=0.4, help="Delay between online translation requests")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = args.input.expanduser().resolve()
    output_path = args.output.expanduser().resolve()
    manifests = []

    if args.batch or input_path.is_dir():
        output_path.mkdir(parents=True, exist_ok=True)
        cache_dir = output_path / "_translation_cache"
        for pdf_path in iter_pdfs(input_path):
            docx_path = output_path / f"{safe_stem(pdf_path)}.docx"
            manifests.append(convert_one(pdf_path, docx_path, args, cache_dir))
    else:
        cache_dir = output_path.parent / "_translation_cache"
        manifests.append(convert_one(input_path, output_path, args, cache_dir))

    summary = {
        "input": str(input_path),
        "output": str(output_path),
        "file_count": len(manifests),
        "total_pages": sum(m["page_count"] for m in manifests),
        "total_ocr_pages": sum(m["ocr_pages"] for m in manifests),
        "total_translated_pages": sum(m["translated_pages"] for m in manifests),
        "files": manifests,
    }
    summary_path = output_path / "_pdf_to_chinese_docx_summary.json" if output_path.is_dir() else output_path.with_suffix(".summary.json")
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
