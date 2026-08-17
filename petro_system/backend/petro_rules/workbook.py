from __future__ import annotations

import hashlib
import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable, Optional, Union

from openpyxl import load_workbook


HEADERS = (
    "序号",
    "规则编号",
    "规则名称",
    "规则类别",
    "知识类型",
    "子类型",
    "适用范围",
    "触发条件",
    "规则内容",
    "参数范围",
    "来源文件",
    "来源页码",
    "来源依据",
    "备注",
)


@dataclass(frozen=True)
class CategorySpec:
    code: str
    sheet_name: str
    knowledge_type: str
    expected_count: int
    rule_prefix: str


CATEGORY_SPECS = (
    CategorySpec("SOP", "操作规程规则", "操作规程类", 257, "SOP-"),
    CategorySpec("QS", "质量标准规则", "质量标准类", 158, "QS-"),
    CategorySpec("FLOW", "流程特征规则", "流程特征类", 135, "FLOW-"),
    CategorySpec("PROC", "工艺规范规则", "工艺规范类", 275, "PROC-"),
    CategorySpec("EXP", "调度经验规则", "调度经验类", 225, "EXP-"),
)


@dataclass(frozen=True)
class RuleRecord:
    sequence_no: int
    rule_no: str
    rule_name: str
    category_code: str
    category_name: str
    knowledge_type: str
    subtype: str
    applicable_scope: str
    trigger_condition: Optional[str]
    rule_content: str
    parameter_text: Optional[str]
    source_file: str
    source_page: str
    source_basis: str
    notes: Optional[str]
    source_sheet: str
    source_row: int
    content_hash: str

    def business_values(self) -> dict[str, Any]:
        values = asdict(self)
        values.pop("content_hash")
        return values


@dataclass(frozen=True)
class WorkbookData:
    path: Path
    file_sha256: str
    records: tuple[RuleRecord, ...]
    category_counts: dict[str, int]


class WorkbookValidationError(ValueError):
    pass


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _optional(value: Any) -> Optional[str]:
    text = _text(value)
    return text or None


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _content_hash(values: Iterable[Any]) -> str:
    payload = json.dumps(list(values), ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def read_official_workbook(path: Union[str, Path]) -> WorkbookData:
    workbook_path = Path(path).expanduser().resolve()
    if not workbook_path.is_file():
        raise WorkbookValidationError(f"Excel 文件不存在：{workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    errors: list[str] = []
    records: list[RuleRecord] = []
    seen_rule_numbers: dict[str, str] = {}
    category_counts: dict[str, int] = {}

    try:
        for spec in CATEGORY_SPECS:
            if spec.sheet_name not in workbook.sheetnames:
                errors.append(f"缺少工作表：{spec.sheet_name}")
                continue

            sheet = workbook[spec.sheet_name]
            header = tuple(_text(cell.value) for cell in sheet[1])
            if header != HEADERS:
                errors.append(
                    f"{spec.sheet_name} 列名不一致：应为 {list(HEADERS)}，实际为 {list(header)}"
                )
                continue

            sheet_records: list[RuleRecord] = []
            for row_number, values in enumerate(
                sheet.iter_rows(min_row=2, values_only=True), start=2
            ):
                if not any(_text(value) for value in values):
                    continue
                normalized = [_text(value) for value in values]
                if len(normalized) != len(HEADERS):
                    errors.append(f"{spec.sheet_name} 第 {row_number} 行列数不是 14")
                    continue

                try:
                    sequence_no = int(normalized[0])
                except ValueError:
                    errors.append(
                        f"{spec.sheet_name} 第 {row_number} 行序号不是整数：{normalized[0]}"
                    )
                    continue

                required_indexes = (1, 2, 3, 4, 5, 6, 8, 10, 11, 12)
                missing = [HEADERS[index] for index in required_indexes if not normalized[index]]
                if missing:
                    errors.append(
                        f"{spec.sheet_name} 第 {row_number} 行必填字段为空：{', '.join(missing)}"
                    )
                    continue

                rule_no = normalized[1]
                if rule_no in seen_rule_numbers:
                    errors.append(
                        f"规则编号重复：{rule_no}，位于 {seen_rule_numbers[rule_no]} 和 "
                        f"{spec.sheet_name}!A{row_number}"
                    )
                    continue
                seen_rule_numbers[rule_no] = f"{spec.sheet_name}!A{row_number}"

                if not rule_no.startswith(spec.rule_prefix):
                    errors.append(
                        f"{spec.sheet_name} 第 {row_number} 行规则编号前缀错误：{rule_no}"
                    )
                if normalized[3] != spec.sheet_name:
                    errors.append(
                        f"{spec.sheet_name} 第 {row_number} 行规则类别错误：{normalized[3]}"
                    )
                if normalized[4] != spec.knowledge_type:
                    errors.append(
                        f"{spec.sheet_name} 第 {row_number} 行知识类型错误：{normalized[4]}"
                    )

                hash_values = normalized[1:]
                sheet_records.append(
                    RuleRecord(
                        sequence_no=sequence_no,
                        rule_no=rule_no,
                        rule_name=normalized[2],
                        category_code=spec.code,
                        category_name=normalized[3],
                        knowledge_type=normalized[4],
                        subtype=normalized[5],
                        applicable_scope=normalized[6],
                        trigger_condition=_optional(values[7]),
                        rule_content=normalized[8],
                        parameter_text=_optional(values[9]),
                        source_file=normalized[10],
                        source_page=normalized[11],
                        source_basis=normalized[12],
                        notes=_optional(values[13]),
                        source_sheet=spec.sheet_name,
                        source_row=row_number,
                        content_hash=_content_hash(hash_values),
                    )
                )

            sequences = [record.sequence_no for record in sheet_records]
            expected_sequences = list(range(1, len(sheet_records) + 1))
            if sequences != expected_sequences:
                errors.append(f"{spec.sheet_name} 序号不是从 1 开始的连续序列")
            if len(sheet_records) != spec.expected_count:
                errors.append(
                    f"{spec.sheet_name} 应有 {spec.expected_count} 条，实际 {len(sheet_records)} 条"
                )
            category_counts[spec.sheet_name] = len(sheet_records)
            records.extend(sheet_records)
    finally:
        workbook.close()

    expected_total = sum(spec.expected_count for spec in CATEGORY_SPECS)
    if len(records) != expected_total:
        errors.append(f"五类规则应共 {expected_total} 条，实际 {len(records)} 条")
    if errors:
        preview = "\n".join(f"- {error}" for error in errors[:30])
        suffix = f"\n- 其余 {len(errors) - 30} 个错误已省略" if len(errors) > 30 else ""
        raise WorkbookValidationError(f"Excel 校验失败：\n{preview}{suffix}")

    return WorkbookData(
        path=workbook_path,
        file_sha256=_file_sha256(workbook_path),
        records=tuple(records),
        category_counts=category_counts,
    )
