"""Extract source-backed, non-duplicate tank rules for the 2026.9.29 snapshot."""

import json
import re
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source-materials/20260727调研资料"
RESULT = ROOT / "extract_result/2026.9.29"
TANK_RESULT = RESULT / "罐区类"


def tank_key(plant, tank):
    return plant, str(tank).strip().removeprefix("T-")


def rule(kind, plant, tank, material, value):
    suffix = re.sub(r"[^A-Za-z0-9]", "", tank).upper()
    plant_id = "YL" if plant == "延炼" else "YSH"
    label = "下限" if kind == "MIN" else "上限"
    param = "inventory_level_min" if kind == "MIN" else "inventory_level_max"
    return {
        "rule_id": f"TANK_LEVEL_{kind}_{plant_id}_{suffix}",
        "description": f"{plant}{tank}罐（{material}）安全液位{label}为{value:g}米。",
        "sets": [{"set_name": "T", "subset_name": None, "element": tank}],
        "parameters": {param: {"value": value, "unit": "m"}},
    }


def main():
    crude = {
        x["sets"][0]["element"]
        for x in json.loads((RESULT / "原油类/原油罐边界.json").read_text())
    }
    workbook_0724 = openpyxl.load_workbook(
        SOURCE / "延炼延石化_储罐与装置侧线与原油_数据自整理_0724.xlsx",
        read_only=True,
        data_only=True,
    )
    workbook_0720 = openpyxl.load_workbook(
        SOURCE / "延炼延石化数据收集_0720.xlsx", read_only=True, data_only=True
    )
    source_0724 = workbook_0724["储罐数据"]
    source_0720 = workbook_0720["01_储罐基础数据"]

    primary = {}
    partial_upper = {}
    for row_number, row in enumerate(source_0724.iter_rows(min_row=5, values_only=True), 5):
        for plant, offset in (("延炼", 0), ("延石化", 20)):
            material, tank, _, high, low, *rest = row[offset : offset + 10]
            if tank and isinstance(high, (int, float)) and not isinstance(low, (int, float)):
                partial_upper[tank_key(plant, tank)] = (high, row_number)
            if not (tank and isinstance(high, (int, float)) and isinstance(low, (int, float))):
                continue
            tank = str(tank).strip()
            if "待" in tank or "未知" in tank:
                continue
            assert 0 <= low < high, (row_number, plant, tank, high, low)
            key = tank_key(plant, tank)
            assert key not in primary, key
            primary[key] = (plant, tank, material, high, low, rest[-1], row_number)

    collected = {}
    for row_number, row in enumerate(source_0720.iter_rows(min_row=6, values_only=True), 6):
        plant = "延炼" if "延炼" in str(row[1]) else "延石化"
        tank_group, high, low = row[4], row[10], row[11]
        if not (tank_group and isinstance(high, (int, float)) and isinstance(low, (int, float))):
            continue
        for tank in re.split(r"[、，,\s]+", str(tank_group).strip()):
            if not tank or "待" in tank:
                continue
            key = tank_key(plant, tank)
            assert key not in collected, (key, row_number)
            collected[key] = (plant, tank, row[3], high, low, row[7], row_number)

    bounds = []
    withheld = []
    for key, (plant, tank, material, high, low, height, row_number) in primary.items():
        if tank in crude:
            continue
        bounds.append(rule("MIN", plant, tank, material, low))
        if plant == "延石化" and tank.startswith(("T-41", "T-42", "T-43", "T-48")):
            withheld.append((tank, "车间罐容表与汇总表上限口径不一致"))
        elif isinstance(height, (int, float)) and high > height:
            withheld.append((tank, f"上限{high:g}米高于罐高{height:g}米"))
        else:
            assert key not in collected or collected[key][3:5] == (high, low), key
            bounds.append(rule("MAX", plant, tank, material, high))

    for key, (plant, tank, material, high, low, height, row_number) in collected.items():
        if key in primary or tank in crude:
            continue
        assert plant == "延石化" and 0 <= low < high, (key, high, low)
        assert not isinstance(height, (int, float)) or high <= height, key
        bounds.append(rule("MIN", plant, tank, material, low))
        if key in partial_upper and partial_upper[key][0] != high:
            withheld.append((tank, "0724表仅填写的上限与0720表不一致"))
        else:
            bounds.append(rule("MAX", plant, tank, material, high))

    conditions = []
    for tank in ("T-4201", "T-4202", "T-4203", "T-4204"):
        conditions.append(
            {
                "rule_id": f"TANK_GAS_FRACTIONATION_PUMP_START_PRESSURE_MIN_{tank.replace('-', '')}",
                "description": (
                    f"2026年检修后罐区开工时，{tank}气分原料罐付料并启动P-4202/AB泵前，"
                    "罐压力须保持在0.65兆帕以上。"
                ),
                "sets": [{"set_name": "T", "subset_name": None, "element": tank}],
                "parameters": {"pump_start_tank_pressure_min": {"value": 0.65, "unit": "MPa"}},
            }
        )

    assert len(primary) == 158 and len(collected) == 117, (len(primary), len(collected))
    assert len(bounds) == 267 and len(conditions) == 4, (len(bounds), len(conditions))
    assert len(withheld) == 49, len(withheld)
    existing_ids = {
        x["rule_id"]
        for file in RESULT.rglob("*.json")
        if file.name != "规则分类统计.json" and file.parent != TANK_RESULT
        for x in json.loads(file.read_text())
    }
    ids = [x["rule_id"] for x in bounds + conditions]
    assert len(ids) == len(set(ids)) and not (set(ids) & existing_ids)

    for name, data in (("安全液位上下限.json", bounds), ("付油前置条件.json", conditions)):
        (TANK_RESULT / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n")

    stats_file = RESULT / "规则分类统计.json"
    stats = [x for x in json.loads(stats_file.read_text()) if x["小类"] not in ("安全液位上下限", "付油前置条件")]
    insert_at = next(i for i, x in enumerate(stats) if x["大类"] == "罐区类")
    stats[insert_at:insert_at] = [
        {"大类": "罐区类", "小类": "安全液位上下限", "规则数": len(bounds)},
        {"大类": "罐区类", "小类": "付油前置条件", "规则数": len(conditions)},
    ]
    stats_file.write_text(json.dumps(stats, ensure_ascii=False, indent=2) + "\n")
    print(f"Extracted {len(bounds)} level boundaries, {len(conditions)} startup conditions; withheld {len(withheld)} upper bounds.")


if __name__ == "__main__":
    main()
