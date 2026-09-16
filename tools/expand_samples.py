import json
import re
from pathlib import Path
from collections import defaultdict

from openpyxl import load_workbook

ROOT = Path('/Users/machong/shihua-project')
OUT = ROOT / 'extract_result_v2'
SRC = ROOT / 'source-materials' / '20260727调研资料'


def clean(v):
    if isinstance(v, str):
        return re.sub(r'<[^>]+>', '', v).strip()
    return v


def num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(',', '')
    try:
        return float(s)
    except Exception:
        return None


def base(rid, typ, cat, sub, name, desc, subject, pred, obj, source, scope=None,
         record_type='fact', when=None, then=None, mapping=None, parameters=None, sets=None):
    if scope is None:
        scope = {}
    if sets is None:
        sets = [{'set_name': k, 'subset_name': None, 'element': v}
                for k, vals in scope.items() for v in (vals if isinstance(vals, list) else [vals])]
    if parameters is None:
        parameters = {}
        if isinstance(obj, dict):
            for k, v in obj.items():
                if isinstance(v, (int, float, str)):
                    parameters[k] = {'value': v, 'unit': None}
    return {
        'rule_id': rid, 'type': typ, 'category': cat, 'subcategory': sub,
        'name': name, 'description': desc, 'scope': scope,
        'record_type': record_type, 'when': when, 'then': then,
        'mapping': mapping or [], 'subject': subject, 'predicate': pred,
        'object': obj, 'sets': sets, 'parameters': parameters,
        'source': source, 'review_status': 'draft'
    }


def device_subcategory(name):
    s = str(name)
    if '催化' in s:
        return '催化'
    if '重整' in s:
        return '重整'
    if '乙烯' in s:
        return '乙烯'
    if '聚丙烯' in s or '聚烯' in s:
        return '聚烯烃'
    if 'MTBE' in s:
        return 'MTBE'
    if '气体分馏' in s or '气分' in s:
        return '气分'
    if '航煤' in s:
        return '航煤加氢'
    if '柴油' in s:
        return '汽柴油加氢'
    if '苯' in s or '芳烃' in s:
        return '芳烃'
    if '常压' in s or '蒸馏' in s:
        return '一次装置常减压'
    if '焦化' in s:
        return '延迟焦化'
    if '渣油' in s:
        return '渣加'
    if '轻烃' in s:
        return '轻烃回收'
    return '一次装置常减压'


def tank_subcategory(material):
    s = str(material or '')
    if '原油' in s:
        return '原油罐'
    if '渣油' in s:
        return '渣油罐'
    if '柴油' in s:
        return '柴油罐'
    if '汽油' in s or '石脑油' in s:
        return '汽油罐'
    if '航煤' in s or '喷气燃料' in s:
        return '煤油罐'
    if any(k in s for k in ('轻烃', '液化气', '丙烷', '丙烯', 'C4', '碳四')):
        return '轻烃罐'
    if any(k in s for k in ('芳烃', '苯', '二甲苯', '乙苯', '苯乙烯')):
        return '芳烃罐'
    if '蜡油' in s:
        return '蜡油罐'
    if any(k in s for k in ('MTBE', '烷基化', '甲醇', '醋酸', '烃化', '脱氢', '酯后')):
        return '中间产品罐'
    return '中间产品罐'


def add_device_capacity(rules):
    p = SRC / '延炼延石化_储罐与装置侧线与原油_数据自整理_0724.xlsx'
    ws = load_workbook(p, read_only=True, data_only=True)['二次加工装置基础数据']
    for row_no, row in enumerate(ws.iter_rows(min_row=3, values_only=True), 3):
        name = clean(row[1])
        if not name:
            continue
        initial, rated, maximum, max_factor, min_factor = map(num, row[3:8])
        if rated is None or maximum is None:
            continue
        sub = device_subcategory(name)
        obj = {'initial_load': initial, 'rated_load': rated, 'max_load': maximum,
               'max_load_factor': max_factor, 'min_load_factor': min_factor,
               'unit': 't/d'}
        then = {'variable': 'feed_rate', 'operator': 'between',
                'value': [round(rated * (min_factor or 0), 6), maximum], 'unit': 't/d'}
        rid = f'R-UNIT-CAP-{row_no:03d}'
        rules.append(base(rid, 'process_specification', '装置类', sub,
                          f'{name}负荷边界', f'{name}初始负荷{initial}吨/天、额定负荷{rated}吨/天、最大加工负荷{maximum}吨/天。',
                          name, 'feed_rate', obj,
                          {'document': str(p.relative_to(ROOT)), 'sheet': ws.title, 'row': row_no},
                          scope={'unit': [name]}, record_type='rule', then=then,
                          mapping=[{'target': 'MILP', 'type': 'variable_bound'}]))


def add_tanks(rules):
    p = SRC / '延炼延石化_储罐与装置侧线与原油_数据自整理_0724.xlsx'
    ws = load_workbook(p, read_only=True, data_only=True)['储罐数据']
    for row_no, row in enumerate(ws.iter_rows(min_row=5, max_row=122, values_only=True), 5):
        material, tank = clean(row[0]), clean(row[1])
        volume = num(row[2])
        if not material or not tank or volume is None or str(material).startswith('共计'):
            continue
        sub = tank_subcategory(material)
        obj = {'material': material, 'volume': volume, 'unit': 'm3'}
        upper, lower, level, coef, density = map(num, (row[3], row[4], row[6], row[7], row[8]))
        if upper is not None:
            obj.update({'safe_upper_level': upper, 'safe_lower_level': lower,
                        'initial_level': level, 'meter_coefficient': coef,
                        'density': density})
            then = {'variable': 'liquid_level', 'operator': 'between',
                    'value': [lower, upper], 'unit': 'm'}
            typ, rt, mapping = 'process_specification', 'rule', [{'target': 'MILP', 'type': 'variable_bound'}]
        else:
            then, typ, rt, mapping = None, 'process_specification', 'parameter', []
        rid = f'R-TANK-PARAM-{row_no:03d}'
        rules.append(base(rid, typ, '罐区类', sub, f'{tank}罐容与安全液位',
                          f'{material}{tank}罐容为{volume}立方米，来源表中记录了罐容及可用液位参数。',
                          str(tank), 'tank_capacity_and_level', obj,
                          {'document': str(p.relative_to(ROOT)), 'sheet': ws.title, 'row': row_no},
                          scope={'tank': [str(tank)], 'material': [material]}, record_type=rt,
                          then=then, mapping=mapping))


def add_crude_properties(rules):
    p = SRC / '延炼延石化_储罐与装置侧线与原油_数据自整理_0724.xlsx'
    ws = load_workbook(p, read_only=True, data_only=True)['原油性质数据']
    for row_no, row in enumerate(ws.iter_rows(min_row=3, values_only=True), 3):
        cid, name = clean(row[0]), clean(row[1])
        if not cid or not name:
            continue
        api, gravity, sulfur, acid = map(num, row[2:6])
        obj = {'crude_id': cid, 'api': api, 'specific_gravity': gravity,
               'sulfur': sulfur, 'acid_value': acid}
        rules.append(base(f'R-CRUDE-PROP-{row_no:03d}', 'process_specification', '原油类', '原油混合性质计算',
                          f'{name}原油性质参数', f'{name}（{cid}）的API度、比重、硫含量和酸值见原油性质数据表。',
                          name, 'crude_properties', obj,
                          {'document': str(p.relative_to(ROOT)), 'sheet': ws.title, 'row': row_no},
                          scope={'material': [name]}, record_type='fact'))


def add_assay(rules):
    p = SRC / '延炼延石化_储罐与装置侧线与原油_数据自整理_0724.xlsx'
    ws = load_workbook(p, read_only=True, data_only=True)['Assay表']
    section = None
    section_sub = None
    for row_no, row in enumerate(ws.iter_rows(values_only=True), 1):
        marker = clean(row[0])
        if marker and str(marker).startswith('*'):
            section = str(marker).lstrip('*')
            section_sub = '原油混合侧线产品收率计算' if '质量收率' in section else '原油混合性质计算'
            continue
        if not marker or not clean(row[1]) or section is None:
            continue
        stream_id, stream_name = marker, clean(row[1])
        for col, crude in ((2, 'YCO'), (3, 'RCO')):
            value = num(row[col]) if col < len(row) else None
            if value is None:
                continue
            obj = {'crude_id': crude, 'side_stream_id': stream_id,
                   'side_stream': stream_name, 'value': value, 'unit': None}
            rules.append(base(f'R-CRUDE-ASSAY-{row_no:03d}-{col}', 'process_specification', '原油类', section_sub,
                              f'{crude}-{stream_id}-{section}', f'{crude}加工时{stream_name}的{section}为{value}。',
                              stream_id, section, obj,
                              {'document': str(p.relative_to(ROOT)), 'sheet': ws.title, 'row': row_no, 'column': col + 1},
                              scope={'material': [stream_name], 'crude': [crude]}, record_type='parameter'))


def point_subcategory(sheet):
    s = str(sheet)
    if '催化' in s: return '催化'
    if '重整' in s: return '重整'
    if 'MTBE' in s: return 'MTBE'
    if '气体分馏' in s: return '气分'
    if '航煤' in s: return '航煤加氢'
    if '柴油' in s: return '汽柴油加氢'
    if '汽油' in s: return 'S-ZORB'
    if '乙苯' in s or '苯乙烯' in s: return '芳烃'
    return '一次装置常减压'


def add_point_properties(rules):
    paths = [SRC / '延炼装置点位性质全数据_0723.xlsx', SRC / '延石化装置点位性质全数据_0723.xlsx']
    for p in paths:
        wb = load_workbook(p, read_only=True, data_only=True)
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [clean(v) for v in rows[0]]
            sub = point_subcategory(sheet.title)
            for row_no, row in enumerate(rows[1:], 2):
                code, sample, device = (clean(row[i]) if i < len(row) else None for i in (0, 1, 2))
                if not code or not sample:
                    continue
                indicators = {}
                for i, value in enumerate(row[3:], 3):
                    v = num(value)
                    if v is not None and i < len(headers) and headers[i]:
                        indicators[headers[i]] = v
                if len(indicators) < 2:
                    continue
                obj = {'sample_dot_code': code, 'sample_name': sample,
                       'device_name': device, 'indicators': indicators}
                params = {k: {'value': v, 'unit': None} for k, v in indicators.items()}
                rules.append(base(f'R-UNIT-POINT-{len(rules)+1:05d}', 'process_specification', '装置类', sub,
                                  f'{device or sheet.title}-{sample}点位性质', f'{device or sheet.title}的{sample}点位包含{len(indicators)}项数值性质指标。',
                                  sample, 'measured_properties', obj,
                                  {'document': str(p.relative_to(ROOT)), 'sheet': sheet.title, 'row': row_no},
                                  scope={'unit': [device or sheet.title], 'material': [sample]}, record_type='fact', parameters=params))


def main():
    all_path = OUT / 'all_rule_samples.json'
    rules = json.loads(all_path.read_text())['rules']
    existing = {r['rule_id'] for r in rules}
    add_device_capacity(rules)
    add_tanks(rules)
    add_crude_properties(rules)
    add_assay(rules)
    add_point_properties(rules)
    # De-duplicate defensively if a source is re-run.
    unique = []
    seen = set()
    for r in rules:
        if r['rule_id'] not in seen:
            unique.append(r); seen.add(r['rule_id'])
    rules = unique
    meta = json.loads(all_path.read_text()).get('meta', {})
    meta.update({'generated_at': '2026-09-08', 'rule_count': len(rules),
                 'expansion_sources': ['source-materials/20260727调研资料/延炼延石化_储罐与装置侧线与原油_数据自整理_0724.xlsx',
                                       'source-materials/20260727调研资料/延炼装置点位性质全数据_0723.xlsx',
                                       'source-materials/20260727调研资料/延石化装置点位性质全数据_0723.xlsx'],
                 'sample_policy': {'target_per_subcategory': '按原始数据密度扩充；数值快照保留为fact/parameter'}})
    all_path.write_text(json.dumps({'meta': meta, 'rules': rules}, ensure_ascii=False, indent=2))
    by = defaultdict(list)
    for r in rules:
        by[(r['category'], r['subcategory'])].append(r)
    for catdir in [p for p in OUT.iterdir() if p.is_dir()]:
        for f in catdir.glob('*.json'):
            data = json.loads(f.read_text())
            rs = by.get((data['category'], data['subcategory']), [])
            data['rules'] = rs; data['sample_count'] = len(rs)
            data['status'] = 'ready' if rs else 'pending'
            data['notes'] = '' if rs else '待补充可靠原始资料样例。'
            f.write_text(json.dumps(data, ensure_ascii=False, indent=2))
    idx = json.loads((OUT / 'index.json').read_text())
    idx['rule_count'] = len(rules)
    for c in idx['categories']:
        c['rule_count'] = 0
        for item in c['files']:
            rs = by.get((c['category'], item['subcategory']), [])
            item['sample_count'] = len(rs); item['status'] = 'ready' if rs else 'pending'
            c['rule_count'] += len(rs)
    (OUT / 'index.json').write_text(json.dumps(idx, ensure_ascii=False, indent=2))
    report = json.loads((OUT / 'extraction_report.json').read_text())
    report['rule_count'] = len(rules)
    report['expanded_categories'] = {c['category']: c['rule_count'] for c in idx['categories']}
    report['record_type_counts'] = defaultdict(int)
    for r in rules: report['record_type_counts'][r['record_type']] += 1
    (OUT / 'extraction_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print('expanded', len(rules), 'from', len(existing))


if __name__ == '__main__':
    main()
