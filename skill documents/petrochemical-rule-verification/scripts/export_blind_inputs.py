#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path
from openpyxl import Workbook
from common import HIDDEN_FROM_AGENTS, header_map, is_rule_sheet, load_xlsx, save_xlsx, safe_str

AGENT1_COLUMNS = [
    "规则编号", "规则名称", "知识类型", "子类型", "适用范围",
    "规则内容mask1", "关键字段mask1", "填空内容mask1", "原文内容mask1", "来源文件", "具体页码",
]
AGENT2_COLUMNS = [
    "规则编号", "规则名称", "知识类型", "子类型", "适用范围",
    "规则内容mask2", "关键字段mask2", "填空内容mask2", "原文内容mask2", "来源文件", "具体页码",
]


def copy_visible_sheet(src_ws, dst_ws, cols):
    h = header_map(src_ws)
    dst_ws.append(cols)
    for row in range(2, src_ws.max_row + 1):
        if not safe_str(src_ws.cell(row, h.get("规则编号", 1)).value):
            continue
        dst_ws.append([src_ws.cell(row, h[col]).value if col in h else None for col in cols])
    for col_idx, name in enumerate(cols, 1):
        dst_ws.cell(1, col_idx).value = name
        dst_ws.column_dimensions[dst_ws.cell(1, col_idx).column_letter].width = 24 if "内容" in name else 16


def create_blind_inputs(input_path: str, outdir: str):
    base = load_xlsx(input_path)
    Path(outdir).mkdir(parents=True, exist_ok=True)
    for agent_name, cols, filename in [
        ("agent1", AGENT1_COLUMNS, "agent1_blind_input.xlsx"),
        ("agent2", AGENT2_COLUMNS, "agent2_blind_input.xlsx"),
    ]:
        wb = Workbook()
        default = wb.active
        wb.remove(default)
        for src_ws in base.worksheets:
            if not is_rule_sheet(src_ws):
                continue
            dst_ws = wb.create_sheet(src_ws.title)
            # Hard guard: do not export hidden columns even if caller modifies column list.
            final_cols = [c for c in cols if c not in HIDDEN_FROM_AGENTS]
            copy_visible_sheet(src_ws, dst_ws, final_cols)
        save_xlsx(wb, str(Path(outdir) / filename))


def main():
    parser = argparse.ArgumentParser(description="Export blind input workbooks for two independent verification agents.")
    parser.add_argument("--input", required=True, help="Input mask workbook")
    parser.add_argument("--outdir", required=True, help="Output directory")
    args = parser.parse_args()
    create_blind_inputs(args.input, args.outdir)


if __name__ == "__main__":
    main()
