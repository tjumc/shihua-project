#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

STANDARD_COLUMNS = [
    "规则编号",
    "规则名称",
    "知识类型",
    "子类型",
    "适用范围",
    "规则内容",
    "来源文件",
    "具体页码",
    "原文内容",
    "JSON",
    "备注",
]

DEFAULT_MATCH_FIELDS = ["知识类型", "子类型", "适用范围", "规则名称"]
DEFAULT_CONTENT_FIELDS = ["规则内容", "来源文件", "具体页码", "原文内容", "JSON", "备注"]

PROCESS_TYPE = "工艺规范类"
EXPERIENCE_TYPE = "调度经验类"

EXP_PREFIX_BY_SUBTYPE = {
    "物性平滑": "EXP-SMO",
    "负荷联动": "EXP-LNK",
    "操作频次": "EXP-FRQ",
    "安全边际": "EXP-SAF",
    "库位平衡": "EXP-TNK",
    "预防性腾挪": "EXP-MOV",
    "效益优先": "EXP-BEN",
    "能耗成本": "EXP-ENE",
    "质量卡边": "EXP-QUA",
    "故障导向": "EXP-FLT",
    "指标超差": "EXP-DEV",
    "极端天气": "EXP-WEA",
    "外部中断": "EXP-WEA",
    "其他": "EXP-OTH",
}


def safe_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize(value: Any) -> str:
    s = safe_str(value)
    replacements = {
        "（": "(",
        "）": ")",
        "，": ",",
        "。": ".",
        "；": ";",
        "：": ":",
        "％": "%",
        "　": " ",
        "～": "~",
        "—": "-",
        "－": "-",
    }
    for old, new in replacements.items():
        s = s.replace(old, new)
    s = re.sub(r"\s+", "", s)
    return s.lower()


def header_map(ws: Worksheet) -> Dict[str, int]:
    out: Dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(1, col).value
        if val is not None:
            out[safe_str(val)] = col
    return out


def is_rule_sheet(ws: Worksheet) -> bool:
    h = header_map(ws)
    return "规则编号" in h and "规则内容" in h


def sheet_kind(ws: Worksheet, rows: List[Dict[str, Any]] | None = None) -> str:
    title = ws.title
    if "工艺" in title or "规范" in title:
        return PROCESS_TYPE
    if "调度" in title or "经验" in title:
        return EXPERIENCE_TYPE
    if rows:
        counts = Counter(safe_str(r.get("知识类型")) for r in rows)
        if counts:
            return counts.most_common(1)[0][0] or title
    return title


def row_to_dict(ws: Worksheet, row: int) -> Dict[str, Any]:
    h = header_map(ws)
    return {col: ws.cell(row, h[col]).value if col in h else None for col in STANDARD_COLUMNS}


def load_rule_rows(path: Path) -> Dict[str, List[Dict[str, Any]]]:
    wb = load_workbook(path, data_only=False)
    grouped: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for ws in wb.worksheets:
        if not is_rule_sheet(ws):
            continue
        rows: List[Dict[str, Any]] = []
        h = header_map(ws)
        for row in range(2, ws.max_row + 1):
            if not safe_str(ws.cell(row, h["规则编号"]).value) and not safe_str(ws.cell(row, h["规则内容"]).value):
                continue
            record = row_to_dict(ws, row)
            record["_source_sheet"] = ws.title
            rows.append(record)
        grouped[sheet_kind(ws, rows)].extend(rows)
    return dict(grouped)


def make_key(row: Dict[str, Any], fields: Iterable[str]) -> Tuple[str, ...]:
    return tuple(normalize(row.get(field)) for field in fields)


def comparable_changed(old: Dict[str, Any], new: Dict[str, Any], fields: Iterable[str]) -> List[str]:
    changed = []
    for field in fields:
        if normalize(old.get(field)) != normalize(new.get(field)):
            changed.append(field)
    return changed


def serializable_row(row: Dict[str, Any], fields: Iterable[str]) -> Dict[str, str]:
    return {field: safe_str(row.get(field)) for field in fields}


def copy_cell_style(src, dst) -> None:
    if src.has_style:
        dst._style = copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    if src.alignment:
        dst.alignment = copy(src.alignment)
    if src.fill:
        dst.fill = copy(src.fill)
    if src.font:
        dst.font = copy(src.font)
    if src.border:
        dst.border = copy(src.border)


def ensure_headers(ws: Worksheet) -> None:
    for idx, col in enumerate(STANDARD_COLUMNS, 1):
        ws.cell(1, idx).value = col


def clear_data_rows(ws: Worksheet) -> None:
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def write_rows(ws: Worksheet, rows: List[Dict[str, Any]]) -> None:
    ensure_headers(ws)
    data_styles = {}
    if ws.max_row >= 2:
        for cidx in range(1, len(STANDARD_COLUMNS) + 1):
            data_styles[cidx] = copy(ws.cell(2, cidx)._style)
    for ridx, row in enumerate(rows, 2):
        for cidx, col in enumerate(STANDARD_COLUMNS, 1):
            cell = ws.cell(ridx, cidx)
            cell.value = row.get(col)
            if cidx in data_styles:
                cell._style = copy(data_styles[cidx])


def find_output_sheet(wb, kind: str) -> Worksheet:
    candidates = [ws for ws in wb.worksheets if is_rule_sheet(ws)]
    for ws in candidates:
        if sheet_kind(ws) == kind:
            return ws
    if kind == PROCESS_TYPE:
        return wb.create_sheet("2、工艺规范规则")
    if kind == EXPERIENCE_TYPE:
        return wb.create_sheet("4、调度经验规则")
    return wb.create_sheet(kind[:31])


def exp_prefix(subtype: str) -> str:
    for key, prefix in EXP_PREFIX_BY_SUBTYPE.items():
        if key in subtype:
            return prefix
    return "EXP-OTH"


def renumber_rows(grouped: Dict[str, List[Dict[str, Any]]], keep_old_ids: bool) -> None:
    if keep_old_ids:
        return
    for kind, rows in grouped.items():
        if kind == PROCESS_TYPE:
            for idx, row in enumerate(rows, 1):
                row["规则编号"] = f"R1-{idx:03d}"
        elif kind == EXPERIENCE_TYPE:
            counters: Counter[str] = Counter()
            for row in rows:
                prefix = exp_prefix(safe_str(row.get("子类型")))
                counters[prefix] += 1
                row["规则编号"] = f"{prefix}-{counters[prefix]:03d}"
        else:
            for idx, row in enumerate(rows, 1):
                row["规则编号"] = f"RULE-{idx:03d}"


def merge_group(
    kind: str,
    old_rows: List[Dict[str, Any]],
    new_rows: List[Dict[str, Any]],
    match_fields: List[str],
    content_fields: List[str],
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], Counter]:
    stats: Counter = Counter()
    log: List[Dict[str, Any]] = []
    merged = [dict(r) for r in old_rows]
    key_to_index: Dict[Tuple[str, ...], int] = {}

    for idx, row in enumerate(merged):
        key = make_key(row, match_fields)
        if key in key_to_index:
            stats["旧表重复"] += 1
            log.append({
                "sheet": kind,
                "action": "旧表重复",
                "rule_id": safe_str(row.get("规则编号")),
                "match_key": " | ".join(key),
                "changed_fields": "",
                "old_values": "",
                "new_values": json.dumps(serializable_row(row, STANDARD_COLUMNS), ensure_ascii=False),
            })
        else:
            key_to_index[key] = idx

    seen_new: set[Tuple[str, ...]] = set()
    for new_row in new_rows:
        key = make_key(new_row, match_fields)
        if key in seen_new:
            stats["新表重复覆盖"] += 1
            log.append({
                "sheet": kind,
                "action": "新表重复覆盖",
                "rule_id": safe_str(new_row.get("规则编号")),
                "match_key": " | ".join(key),
                "changed_fields": "",
                "old_values": "",
                "new_values": json.dumps(serializable_row(new_row, STANDARD_COLUMNS), ensure_ascii=False),
            })
        seen_new.add(key)

        if key not in key_to_index:
            merged.append(dict(new_row))
            key_to_index[key] = len(merged) - 1
            stats["新增"] += 1
            log.append({
                "sheet": kind,
                "action": "新增",
                "rule_id": safe_str(new_row.get("规则编号")),
                "match_key": " | ".join(key),
                "changed_fields": ",".join(STANDARD_COLUMNS),
                "old_values": "",
                "new_values": json.dumps(serializable_row(new_row, STANDARD_COLUMNS), ensure_ascii=False),
            })
            continue

        old_idx = key_to_index[key]
        old_row = merged[old_idx]
        changed = comparable_changed(old_row, new_row, content_fields)
        if changed:
            merged[old_idx] = dict(new_row)
            stats["更新"] += 1
            log.append({
                "sheet": kind,
                "action": "更新",
                "rule_id": safe_str(old_row.get("规则编号")),
                "match_key": " | ".join(key),
                "changed_fields": ",".join(changed),
                "old_values": json.dumps(serializable_row(old_row, changed), ensure_ascii=False),
                "new_values": json.dumps(serializable_row(new_row, changed), ensure_ascii=False),
            })
        else:
            stats["不变"] += 1

    stats["旧表保留"] = len(old_rows) - stats["更新"]
    stats["合并后总数"] = len(merged)
    return merged, log, stats


def write_change_logs(outdir: Path, logs: List[Dict[str, Any]]) -> None:
    fields = ["sheet", "action", "rule_id", "match_key", "changed_fields", "old_values", "new_values"]
    with (outdir / "change_log.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(logs)
    (outdir / "change_log.json").write_text(json.dumps(logs, ensure_ascii=False, indent=2), encoding="utf-8")


def write_report(outdir: Path, old_path: Path, new_path: Path, stats_by_sheet: Dict[str, Counter], match_fields: List[str], content_fields: List[str]) -> None:
    lines = [
        "# 石化规则增量更新报告",
        "",
        f"- 生成时间: {datetime.now().isoformat(timespec='seconds')}",
        f"- 旧总表: {old_path}",
        f"- 新抽取表: {new_path}",
        f"- 匹配字段: {', '.join(match_fields)}",
        f"- 内容比较字段: {', '.join(content_fields)}",
        "",
        "## 汇总",
        "",
    ]
    total = Counter()
    for kind, stats in stats_by_sheet.items():
        total.update(stats)
        lines.append(f"### {kind}")
        for key in ["新增", "更新", "不变", "旧表重复", "新表重复覆盖", "合并后总数"]:
            if stats.get(key):
                lines.append(f"- {key}: {stats[key]}")
        lines.append("")
    lines.extend([
        "### 全部",
        f"- 新增: {total.get('新增', 0)}",
        f"- 更新: {total.get('更新', 0)}",
        f"- 不变: {total.get('不变', 0)}",
        f"- 合并后总数: {total.get('合并后总数', 0)}",
        "",
        "## 下游核验",
        "",
        "将 `extracted_rules.xlsx` 传给 `petrochemical-rule-verification`：",
        "",
        "```bash",
        "python ~/.codex/skills/petrochemical-rule-verification/scripts/run_verification.py \\",
        "  --phase prep \\",
        "  --input extracted_rules.xlsx \\",
        "  --output-dir verification_output",
        "```",
        "",
    ])
    (outdir / "update_report.md").write_text("\n".join(lines), encoding="utf-8")


def merge_workbooks(args: argparse.Namespace) -> None:
    old_path = Path(args.old).expanduser().resolve()
    new_path = Path(args.new).expanduser().resolve()
    outdir = Path(args.output_dir).expanduser().resolve()
    outdir.mkdir(parents=True, exist_ok=True)

    if not old_path.exists():
        raise FileNotFoundError(f"Old workbook not found: {old_path}")
    if not new_path.exists():
        raise FileNotFoundError(f"New workbook not found: {new_path}")

    old_grouped = load_rule_rows(old_path)
    new_grouped = load_rule_rows(new_path)
    all_kinds = list(dict.fromkeys(list(old_grouped.keys()) + list(new_grouped.keys())))

    merged_grouped: Dict[str, List[Dict[str, Any]]] = {}
    all_logs: List[Dict[str, Any]] = []
    stats_by_sheet: Dict[str, Counter] = {}

    for kind in all_kinds:
        merged, logs, stats = merge_group(
            kind,
            old_grouped.get(kind, []),
            new_grouped.get(kind, []),
            args.match_fields,
            args.content_fields,
        )
        merged_grouped[kind] = merged
        all_logs.extend(logs)
        stats_by_sheet[kind] = stats

    renumber_rows(merged_grouped, args.keep_old_ids)

    workbook_path = outdir / "updated_rules.xlsx"
    shutil.copy2(old_path, workbook_path)
    wb = load_workbook(workbook_path)
    for kind, rows in merged_grouped.items():
        ws = find_output_sheet(wb, kind)
        clear_data_rows(ws)
        write_rows(ws, rows)
    wb.save(workbook_path)

    extracted_path = outdir / "extracted_rules.xlsx"
    shutil.copy2(workbook_path, extracted_path)

    write_change_logs(outdir, all_logs)
    write_report(outdir, old_path, new_path, stats_by_sheet, args.match_fields, args.content_fields)

    print(f"Wrote {workbook_path}")
    print(f"Wrote {extracted_path}")
    print(f"Wrote {outdir / 'change_log.csv'}")
    print(f"Wrote {outdir / 'update_report.md'}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Merge an old petrochemical rule workbook with newly extracted rules.")
    parser.add_argument("--old", required=True, help="Previous master workbook")
    parser.add_argument("--new", required=True, help="New extracted rules workbook")
    parser.add_argument("--output-dir", required=True, help="Output directory")
    parser.add_argument("--match-fields", nargs="+", default=DEFAULT_MATCH_FIELDS, help="Fields used to identify the same rule")
    parser.add_argument("--content-fields", nargs="+", default=DEFAULT_CONTENT_FIELDS, help="Fields used to decide whether a matched rule changed")
    parser.add_argument("--keep-old-ids", action="store_true", help="Do not reassign rule IDs after merging")
    return parser.parse_args()


def main() -> None:
    merge_workbooks(parse_args())


if __name__ == "__main__":
    main()
