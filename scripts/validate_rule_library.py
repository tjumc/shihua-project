from __future__ import annotations

import json
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    # Compare everything as strings to avoid Excel type noise.
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    return str(value).strip()


def build_default_paths(root: Path) -> dict[str, Path]:
    # Centralize layout assumptions for validation and report export.
    mc_root = root / "提取结果--mc"
    current_dir = mc_root / "01_当前结果"
    trace_dir = mc_root / "02_过程追溯"
    trace_check_dir = trace_dir / "05_规则追溯核对"
    return {
        "root": root,
        "mc_root": mc_root,
        "main_workbook": current_dir / "主总表.xlsx",
        "final_audit_workbook": current_dir / "终版审计表.xlsx",
        "original_audit_workbook": trace_dir / "01_原始审计" / "提取内容逐项审计.xlsx",
        "json_library": root / "json规则库" / "规则库_当前版.json",
        "output_dir": trace_check_dir / "自动核验",
    }


def load_sheet_records(
    workbook_path: Path, sheet_name: str, required_headers: list[str]
) -> list[dict[str, str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"Missing sheet: {workbook_path.name} / {sheet_name}")

    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = [normalize_cell(cell) for cell in next(rows)]
    except StopIteration as exc:
        workbook.close()
        raise ValueError(f"Empty sheet: {workbook_path.name} / {sheet_name}") from exc

    header_index = {header: index for index, header in enumerate(headers)}
    missing = [header for header in required_headers if header not in header_index]
    if missing:
        workbook.close()
        raise ValueError(
            f"Missing headers in {workbook_path.name} / {sheet_name}: {missing}"
        )

    records: list[dict[str, str]] = []
    for row in rows:
        record = {
            header: normalize_cell(row[index]) if index < len(row) else ""
            for header, index in header_index.items()
        }
        if record.get("规则编号"):
            records.append(record)

    workbook.close()
    return records


def load_json_rules(json_path: Path) -> list[dict[str, Any]]:
    payload = json.loads(json_path.read_text(encoding="utf-8"))
    return payload.get("rules", [])


def detect_suspicious_text(text: str) -> list[str]:
    # Heuristics cover the two corruption patterns seen so far.
    findings: list[str] = []
    if any(token in text for token in ("JSON：", "JSON:", "自然语言：", "自然语言:")):
        findings.append("json_residue")
    if text.count("?") >= 3:
        findings.append("question_mark_corruption")
    return findings


def has_approved_reuse_marker(note: str) -> bool:
    return "同源复用确认" in note


def index_by_rule_id(records: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    return {record["规则编号"]: record for record in records}


def append_issue(bucket: list[dict[str, Any]], issue_type: str, **payload: Any) -> None:
    issue = {"type": issue_type}
    issue.update(payload)
    bucket.append(issue)


def validate_rule_library(root: Path | None = None) -> dict[str, Any]:
    root = (root or Path(__file__).resolve().parents[1]).resolve()
    paths = build_default_paths(root)

    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    manual_review: list[dict[str, Any]] = []

    # Later checks assume the full artifact set is present.
    required_files = {
        "main_workbook": paths["main_workbook"],
        "final_audit_workbook": paths["final_audit_workbook"],
        "original_audit_workbook": paths["original_audit_workbook"],
        "json_library": paths["json_library"],
    }
    for label, file_path in required_files.items():
        if not file_path.exists():
            append_issue(errors, "missing_file", label=label, path=str(file_path))

    if errors:
        return {
            "meta": {
                "generated_at": datetime.now().isoformat(timespec="seconds"),
                "root": str(root),
            },
            "summary": {
                "total_rules": 0,
                "workbook_rule_count": 0,
                "final_delivery_count": 0,
                "final_audit_count": 0,
                "original_audit_count": 0,
                "json_rule_count": 0,
                "error_count": len(errors),
                "warning_count": 0,
                "manual_review_count": 0,
            },
            "errors": errors,
            "warnings": warnings,
            "manual_review": manual_review,
        }

    main_rule_records = load_sheet_records(
        paths["main_workbook"],
        "规则总表",
        ["规则编号", "规则名称", "所属大类", "所属子项", "规则正文", "对应文件名", "对应页码"],
    )
    final_delivery_records = load_sheet_records(
        paths["main_workbook"],
        "最终交付表",
        ["规则编号", "规则名称", "所属大类", "所属子项", "规则内容", "来源文件", "来源页码"],
    )
    final_audit_records = load_sheet_records(
        paths["final_audit_workbook"],
        "终版审计明细",
        ["规则编号", "规则名称", "所属大类", "所属子项", "文件名", "页码"],
    )
    original_audit_records = load_sheet_records(
        paths["original_audit_workbook"],
        "逐项审计",
        ["规则编号", "规则名称", "所属大类", "所属子项", "文件名", "页码"],
    )
    json_rules = load_json_rules(paths["json_library"])

    main_rule_index = index_by_rule_id(main_rule_records)
    final_delivery_index = index_by_rule_id(final_delivery_records)
    final_audit_index = index_by_rule_id(final_audit_records)
    original_audit_index = index_by_rule_id(original_audit_records)
    json_rule_index = {normalize_cell(item.get("rule_id")): item for item in json_rules}

    count_summary = {
        "workbook_rule_count": len(main_rule_records),
        "final_delivery_count": len(final_delivery_records),
        "final_audit_count": len(final_audit_records),
        "original_audit_count": len(original_audit_records),
        "json_rule_count": len(json_rules),
    }
    total_rules = len(final_delivery_records)

    if len(main_rule_records) != total_rules:
        append_issue(
            errors,
            "count_mismatch",
            source="规则总表",
            expected=total_rules,
            actual=len(main_rule_records),
        )
    if len(final_audit_records) != total_rules:
        append_issue(
            errors,
            "count_mismatch",
            source="终版审计明细",
            expected=total_rules,
            actual=len(final_audit_records),
        )
    if len(original_audit_records) != total_rules:
        append_issue(
            errors,
            "count_mismatch",
            source="逐项审计",
            expected=total_rules,
            actual=len(original_audit_records),
        )
    if len(json_rules) != total_rules:
        append_issue(
            errors,
            "count_mismatch",
            source="规则库_当前版.json",
            expected=total_rules,
            actual=len(json_rules),
        )

    # The final delivery sheet is the baseline for ID completeness checks.
    reference_ids = set(final_delivery_index)
    for source_name, index in (
        ("规则总表", main_rule_index),
        ("终版审计明细", final_audit_index),
        ("逐项审计", original_audit_index),
        ("规则库_当前版.json", json_rule_index),
    ):
        missing_ids = sorted(reference_ids - set(index))
        extra_ids = sorted(set(index) - reference_ids)
        if missing_ids:
            append_issue(
                errors, "missing_rule_ids", source=source_name, rule_ids=missing_ids
            )
        if extra_ids:
            append_issue(errors, "extra_rule_ids", source=source_name, rule_ids=extra_ids)

    # These fields should stay aligned across all maintained artifacts.
    compare_specs = [
        (
            "规则名称",
            (
                ("规则总表", main_rule_index, "规则名称"),
                ("最终交付表", final_delivery_index, "规则名称"),
                ("终版审计明细", final_audit_index, "规则名称"),
                ("逐项审计", original_audit_index, "规则名称"),
            ),
        ),
        (
            "所属大类",
            (
                ("规则总表", main_rule_index, "所属大类"),
                ("最终交付表", final_delivery_index, "所属大类"),
                ("终版审计明细", final_audit_index, "所属大类"),
                ("逐项审计", original_audit_index, "所属大类"),
            ),
        ),
        (
            "所属子项",
            (
                ("规则总表", main_rule_index, "所属子项"),
                ("最终交付表", final_delivery_index, "所属子项"),
                ("终版审计明细", final_audit_index, "所属子项"),
                ("逐项审计", original_audit_index, "所属子项"),
            ),
        ),
        (
            "来源文件",
            (
                ("规则总表", main_rule_index, "对应文件名"),
                ("最终交付表", final_delivery_index, "来源文件"),
                ("终版审计明细", final_audit_index, "文件名"),
                ("逐项审计", original_audit_index, "文件名"),
            ),
        ),
        (
            "来源页码",
            (
                ("规则总表", main_rule_index, "对应页码"),
                ("最终交付表", final_delivery_index, "来源页码"),
                ("终版审计明细", final_audit_index, "页码"),
                ("逐项审计", original_audit_index, "页码"),
            ),
        ),
    ]

    for rule_id in sorted(reference_ids):
        for field_name, sources in compare_specs:
            values = {
                source_name: source_index[rule_id].get(column_name, "")
                for source_name, source_index, column_name in sources
            }
            unique_values = {value for value in values.values() if value}
            if len(unique_values) > 1:
                append_issue(
                    errors,
                    "field_mismatch",
                    rule_id=rule_id,
                    field=field_name,
                    values=values,
                )

        workbook_content = main_rule_index[rule_id].get("规则正文", "")
        delivery_content = final_delivery_index[rule_id].get("规则内容", "")
        json_description = normalize_cell(json_rule_index.get(rule_id, {}).get("description"))
        for location, text in (
            ("规则总表.规则正文", workbook_content),
            ("最终交付表.规则内容", delivery_content),
            ("规则库_当前版.json.description", json_description),
        ):
            findings = detect_suspicious_text(text)
            if findings:
                append_issue(
                    warnings,
                    "suspicious_text",
                    rule_id=rule_id,
                    location=location,
                    findings=findings,
                )

    # Duplicate names can be valid reuse, so keep them as warnings.
    name_groups: dict[str, list[str]] = defaultdict(list)
    for record in final_delivery_records:
        name_groups[record["规则名称"]].append(record["规则编号"])
    for rule_name, rule_ids in sorted(name_groups.items()):
        if len(rule_ids) > 1:
            append_issue(
                warnings,
                "duplicate_rule_name",
                rule_name=rule_name,
                rule_ids=sorted(rule_ids),
            )

    source_name_groups: dict[tuple[str, str, str], list[dict[str, str]]] = defaultdict(list)
    for record in final_delivery_records:
        key = (record["规则名称"], record["来源文件"], record["来源页码"])
        source_name_groups[key].append(record)
    for (rule_name, source_file, source_page), records in sorted(source_name_groups.items()):
        if len(records) > 1:
            append_issue(
                manual_review,
                "same_name_same_source_page",
                rule_name=rule_name,
                source_file=source_file,
                source_page=source_page,
                rule_ids=sorted(record["规则编号"] for record in records),
                categories=sorted(
                    {f'{record["所属大类"]} / {record["所属子项"]}' for record in records}
                ),
            )

    content_groups: dict[str, list[dict[str, str]]] = defaultdict(list)
    for record in final_delivery_records:
        content = record["规则内容"]
        if content:
            content_groups[content].append(record)
    for content, records in content_groups.items():
        # Cross-category reuse stays visible, but approved reuse should not block
        # the chain as a pending manual-review item forever.
        if len(records) > 1:
            rule_ids = sorted(record["规则编号"] for record in records)
            categories = sorted(
                {f'{record["所属大类"]} / {record["所属子项"]}' for record in records}
            )
            if len(categories) > 1:
                issue_type = "approved_same_source_reuse"
                issue_bucket = warnings
                if not all(
                    has_approved_reuse_marker(record.get("备注", ""))
                    for record in records
                ):
                    issue_type = "same_content_cross_category"
                    issue_bucket = manual_review
                append_issue(
                    issue_bucket,
                    issue_type,
                    rule_ids=rule_ids,
                    categories=categories,
                    sample_content=content[:120],
                )

    summary = {
        "total_rules": total_rules,
        **count_summary,
        "error_count": len(errors),
        "warning_count": len(warnings),
        "manual_review_count": len(manual_review),
    }
    return {
        "meta": {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "root": str(root),
            "source_of_truth": str(paths["main_workbook"]),
        },
        "summary": summary,
        "errors": errors,
        "warnings": warnings,
        "manual_review": manual_review,
    }


def build_markdown_report(report: dict[str, Any]) -> str:
    # Keep the markdown output ready for direct status reporting.
    lines = [
        "# 规则库自动核验报告",
        "",
        f"- 生成时间：`{report['meta']['generated_at']}`",
        f"- 主源文件：`{report['meta']['source_of_truth']}`",
        "",
        "## 摘要",
        "",
        f"- 规则总数：`{report['summary']['total_rules']}`",
        f"- 结构性错误：`{report['summary']['error_count']}`",
        f"- 警告：`{report['summary']['warning_count']}`",
        f"- 待人工复判：`{report['summary']['manual_review_count']}`",
        "",
        "## 数量闭环",
        "",
        f"- 规则总表：`{report['summary']['workbook_rule_count']}`",
        f"- 最终交付表：`{report['summary']['final_delivery_count']}`",
        f"- 终版审计明细：`{report['summary']['final_audit_count']}`",
        f"- 逐项审计：`{report['summary']['original_audit_count']}`",
        f"- JSON 当前版：`{report['summary']['json_rule_count']}`",
        "",
    ]

    def append_issue_block(title: str, issues: list[dict[str, Any]]) -> None:
        lines.append(f"## {title}")
        lines.append("")
        if not issues:
            lines.append("- 无")
            lines.append("")
            return
        for issue in issues:
            issue_type = issue["type"]
            if issue_type == "same_name_same_source_page":
                lines.append(
                    f"- `{issue['rule_name']}` | 规则：`{' / '.join(issue['rule_ids'])}` | 来源：`{issue['source_file']}` 第 `{issue['source_page']}` 页"
                )
                lines.append(f"  类别：`{'；'.join(issue['categories'])}`")
            elif issue_type == "same_content_cross_category":
                lines.append(
                    f"- 同内容跨类别：`{' / '.join(issue['rule_ids'])}` | 类别：`{'；'.join(issue['categories'])}`"
                )
                lines.append(f"  内容样例：`{issue['sample_content']}`")
            elif issue_type == "approved_same_source_reuse":
                lines.append(
                    f"- 已确认同源复用：`{' / '.join(issue['rule_ids'])}` | 类别：`{'；'.join(issue['categories'])}`"
                )
                lines.append(f"  内容样例：`{issue['sample_content']}`")
            elif issue_type == "duplicate_rule_name":
                lines.append(
                    f"- 重复名称：`{issue['rule_name']}` | 规则：`{' / '.join(issue['rule_ids'])}`"
                )
            elif issue_type == "suspicious_text":
                lines.append(
                    f"- `{issue['rule_id']}` | `{issue['location']}` | 异常：`{' / '.join(issue['findings'])}`"
                )
            else:
                lines.append(f"- `{issue_type}`：`{json.dumps(issue, ensure_ascii=False)}`")
        lines.append("")

    append_issue_block("结构性错误", report["errors"])
    append_issue_block("警告", report["warnings"])
    append_issue_block("待人工复判", report["manual_review"])
    return "\n".join(lines)


def write_validation_report(
    report: dict[str, Any], output_dir: Path
) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = date.today().isoformat()
    json_path = output_dir / f"{stamp}_规则库自动核验报告.json"
    markdown_path = output_dir / f"{stamp}_规则库自动核验报告.md"
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    markdown_path.write_text(build_markdown_report(report), encoding="utf-8")
    return json_path, markdown_path


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    report = validate_rule_library(root)
    output_dir = build_default_paths(root)["output_dir"]
    json_path, markdown_path = write_validation_report(report, output_dir)
    print(f"JSON report: {json_path}")
    print(f"Markdown report: {markdown_path}")
    return 1 if report["errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
